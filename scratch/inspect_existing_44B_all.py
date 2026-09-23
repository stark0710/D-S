import openpyxl

excel_path = "reports/fixed_wing_44B_manual_review.xlsx"
wb = openpyxl.load_workbook(excel_path, read_only=True)
ws = wb["COMPLETE_RESULTS"]

print("Total rows in COMPLETE_RESULTS:", ws.max_row)

print("First 15 rows in COMPLETE_RESULTS:")
for r in range(1, 16):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
    print(f"Row {r}:", row_vals)
