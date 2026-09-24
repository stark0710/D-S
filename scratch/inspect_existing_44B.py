import openpyxl
import pandas as pd

excel_path = "reports/fixed_wing_44B_manual_review.xlsx"

wb = openpyxl.load_workbook(excel_path, read_only=True)
print("Sheet names:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\nSheet {sheet_name} columns:")
    headers = [ws.cell(row=1, column=c).value for c in range(1, 35)]
    headers = [h for h in headers if h is not None]
    print(headers)
    print("Row 2 preview:")
    row2 = [ws.cell(row=2, column=c).value for c in range(1, len(headers) + 1)]
    print(row2)
