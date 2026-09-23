import openpyxl

excel_path = "reports/vehicle_selection_44A_manual_review.xlsx"
wb = openpyxl.load_workbook(excel_path, read_only=True)
ws = wb["MANUAL_REVIEW"]

print("Header:")
header = [ws.cell(row=1, column=c).value for c in range(1, 20)]
print(header)

fw_count = 0
for r in range(2, ws.max_row + 1):
    selected = ws.cell(row=r, column=11).value
    if selected == "FIXED_WING":
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 20)]
        print(f"Row {r} (Case ID {row_vals[0]}):", row_vals[:12])
        fw_count += 1
        if fw_count >= 10:
            break
